import re, json, numpy as np, pandas as pd
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[1]   # project root (script lives in scripts/)
T = ROOT/"tables"; F = ROOT/"figures"; T = ROOT/"tables"; F = ROOT/"figures"
rd = lambda n, **k: pd.read_csv(T/n, **k)
S = json.load(open(T/"results_summary.json"))

# ------------------------------------------------------------ numbers from real outputs
cd = rd("class_distribution.csv", index_col=0)
dq = rd("range_validation.csv", index_col=0)
iqr = rd("iqr_outlier_diagnostic.csv", index_col=0)
desc = rd("descriptive_statistics.csv", index_col=0)
cm_means = rd("class_means.csv", index_col=0)
knn = rd("knn_comparison_table.csv")
pc5 = rd("knn_k5_per_class_metrics.csv", index_col=0)
cv = rd("knn_cross_validation.csv", index_col=0)
cm5 = rd("confusion_matrix_knn_k5.csv", index_col=0)
cm5.index = cm5.index.str.upper(); cm5.columns = cm5.columns.str.upper()
sim = rd("most_confused_pair_feature_similarity.csv", index_col=0)
ct = rd("class_cluster_crosstab_test.csv", index_col=0)
ctr = rd("class_cluster_crosstab_train.csv", index_col=0)
ktab = rd("kmeans_cluster_table_train.csv", index_col=0)
pur = rd("kmeans_purity.csv", index_col=0)
kscan = rd("kmeans_k_scan.csv", index_col=0)
dist = rd("distance_concentration_demo.csv", index_col=0)
hp = rd("high_correlation_pairs.csv")
tt = rd("train_test_class_distribution.csv", index_col=0)
sc = rd("scaling_effect.csv", index_col=0)

n_tot = S["rows"]; n_tr = S["train_rows"]; n_te = S["test_rows"]
big, small = cd.index[0], cd.index[-1]
ratio = S["imbalance_ratio"]
k = {r.K if r.Scale == "Yes" else "raw": r for r in knn.itertuples()}
acc = {kk: knn[(knn.K == kk) & (knn.Scale == "Yes")]["Test Accuracy"].iloc[0] for kk in (3, 5, 9)}
se = float(np.sqrt(acc[5]*(1-acc[5])/n_te))
spread = max(acc.values()) - min(acc.values())
corr_pairs = len(hp)
def cor(a, b):
    r = hp[((hp.iloc[:, 0] == a) & (hp.iloc[:, 1] == b)) | ((hp.iloc[:, 0] == b) & (hp.iloc[:, 1] == a))]
    return float(r.correlation.iloc[0])
r_ap, r_ae, r_pe = cor("Area", "Perimeter"), cor("Area", "EquivDiameter"), cor("Perimeter", "EquivDiameter")
r_ce = cor("Compactness", "Eccentricity")

# raw-distance share of Area+ConvexArea (same recipe as the notebook)
df = pd.read_excel(ROOT/"data/raw/Dry_Bean_Dataset.xlsx").rename(columns={"AspectRation": "AspectRatio", "roundness": "Roundness"})
feats = [c for c in df.columns if c != "Class"]
Xtr, Xte, ytr, yte = train_test_split(df[feats], df["Class"], test_size=0.2, random_state=42, stratify=df["Class"])
rng = np.random.default_rng(42); i, j = rng.integers(0, len(Xtr), 5000), rng.integers(0, len(Xtr), 5000); m = i != j; i, j = i[m], j[m]
sq = (Xtr.to_numpy()[i]-Xtr.to_numpy()[j])**2
share = pd.Series((sq/sq.sum(1, keepdims=True)).mean(0), index=feats)*100
sh_area, sh_conv = share["Area"], share["ConvexArea"]
corr = df[feats].corr(); ev = np.sort(np.linalg.eigvalsh(corr.to_numpy()))[::-1]; cum = np.cumsum(ev/ev.sum())
n90 = int(np.searchsorted(cum, .90)+1); n99 = int(np.searchsorted(cum, .99)+1)
dup_extra = S["duplicates_extra_rows"]
tr_key = set(map(tuple, Xtr.to_numpy())); twins = sum(tuple(r) in tr_key for r in Xte.to_numpy())

classes = list(cm5.index)
cmv = cm5.to_numpy(); diag = np.diag(cmv)
rec5 = pc5["Recall"]; low = rec5.idxmin(); low2 = rec5.drop(low).idxmin()
ab, ba = S["most_confused_pair_errors"]; pair = S["most_confused_pair_k5"]
c2 = int(cm5.loc["BARBUNYA", "CALI"]+cm5.loc["CALI", "BARBUNYA"])
km = S["kmeans"]

# ------------------------------------------------------------ docx helpers
FONT = "DejaVu Sans"; NAVY = RGBColor(0x1F, 0x3A, 0x5F)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6)

def set_font(run, size=10, bold=False, color=None, rtl=True):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    rPr = run._r.get_or_add_rPr(); rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    if rtl:
        rPr.append(OxmlElement("w:rtl"))
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size*2))); rPr.append(szcs)
        if bold: rPr.append(OxmlElement("w:bCs"))
    if color is not None: run.font.color.rgb = color

def rich(p, text, size=10, color=None, rtl=True, bold_all=False):
    for k_, seg in enumerate(re.split(r"\*\*", text)):
        if seg: set_font(p.add_run(seg), size, bold_all or k_ % 2 == 1, color, rtl)

def P(text, size=10, align="both", after=3, before=0, color=None, bold=False, keep=False, indent=0):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.12
    if keep: pf.keep_with_next = True
    if indent: pf.right_indent = Cm(indent)
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    if align == "both": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, text, size, color, True, bold); return p

def H(text):
    p = P(text, size=11.5, align=None, after=3, before=7, color=NAVY, bold=True, keep=True)
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for a, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")): bt.set(qn(a), v)
    b.append(bt); pPr.append(b); return p

def B(text, size=10): return P("•  " + text, size=size, after=2, indent=0.3)

def cell_shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

def table(rows, header, widths, size=8.5, hl_rows=(), first_left=True):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    def fill(cells, vals, bold=False, shade=None):
        for c_, v in zip(cells, vals):
            c_.text = ""; p = c_.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            set_font(p.add_run(str(v)), size, bold, None, rtl=False)
            if shade: cell_shade(c_, shade)
    fill(t.rows[0].cells, header, True, "DCE6F1")
    for n, r in enumerate(rows):
        fill(t.add_row().cells, r, False, "FFF2CC" if n in hl_rows else None)
    t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w*567)))
    for row in t.rows:
        for c_, w in zip(row.cells, widths): c_.width = Cm(w)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2); sp.paragraph_format.line_spacing = 0.5
    return t

def fig(paths, widths, caption):
    t = doc.add_table(rows=1, cols=len(paths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_, pth, w in zip(t.rows[0].cells, paths[::-1], widths[::-1]):   # RTL: first figure sits on the right
        c_.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_.paragraphs[0].paragraph_format.space_after = Pt(0)
        c_.paragraphs[0].add_run().add_picture(str(F/pth), width=Cm(w))
    P(caption, size=8, align="center", after=4, color=RGBColor(0x55, 0x55, 0x55))

# page number footer
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); rr = OxmlElement("w:r"); tt_ = OxmlElement("w:t"); tt_.text = "1"; rr.append(tt_); fld.append(rr); fp._p.append(fld)

f4 = lambda x: f"{x:.4f}"; f3 = lambda x: f"{x:.3f}"
# ------------------------------------------------------------ CONTENT
P("گزارش نهایی پروژه: تشخیص انواع لوبیا با KNN و مقایسهٔ کلاس‌های واقعی با خوشه‌های K-Means", size=15, align="center", after=2, color=NAVY, bold=True)
P(f"دیتاست Dry Bean (UCI) با {n_tot:,} دانه، 16 ویژگی هندسی و 7 کلاس. همهٔ اعداد از اجرای نوت‌بوک به‌دست آمده‌اند و با random_state=42 تکرارپذیرند.", size=8.5, align="center", after=4, color=RGBColor(0x55, 0x55, 0x55))

H("1. مقدمه، هدف و توصیف داده")
P(f"دو پرسش بررسی شد: (1) آیا KNN می‌تواند نوع لوبیا را از روی ویژگی‌های هندسی تشخیص دهد؟ (2) آیا هفت خوشهٔ K-Means که بدون دیدن برچسب ساخته می‌شوند، با هفت کلاس واقعی مطابقت دارند؟ داده شامل **{n_tot:,} نمونه** (هر ردیف یک دانه)، **16 ویژگی عددی** (Area، Perimeter، MajorAxisLength، MinorAxisLength، AspectRatio، Eccentricity، ConvexArea، EquivDiameter، Extent، Solidity، Roundness، Compactness و چهار ShapeFactor) و ستون هدف **Class** با هفت کلاس است: Barbunya، Bombay، Cali، Dermason، Horoz، Seker و Sira. ستون Class فقط هدف است و در هیچ مدلی جزو ورودی نبود (در نوت‌بوک با assert و ممیزی نشت بررسی می‌شود). نام دو ستون فایل خام (AspectRation و roundness) فقط برای یکدستی به AspectRatio و Roundness تغییر یافت.")

H("2. کیفیت داده")
n_inv = int(dq["Invalid rows"].sum())
B(f"**مقدار گمشده:** صفر. **مقدار غیرعددی یا بی‌نهایت در ویژگی‌ها:** صفر. **مقدار خارج از محدودهٔ منطقی:** {n_inv} مورد؛ قاعده‌ها برای 12 ویژگی بررسی شدند (اندازه‌ها > 0، AspectRatio ≥ 1، Eccentricity در [0,1)، و Extent/Solidity/Roundness/Compactness در (0,1]).")
B(f"**ردیف تکراری:** {dup_extra} ردیف تکراری ({dup_extra/n_tot:.2%}) در 68 جفت؛ همهٔ آنها متعلق به کلاس Horoz‌اند و هیچ ویژگی یکسانی برچسب متفاوت ندارد. **تصمیم: حذف نشدند.** دلیل: (الف) پروژه حذف خودکار را نمی‌خواهد و دو دانهٔ متفاوت می‌توانند اندازه‌گیری هندسی یکسان داشته باشند؛ (ب) تکرار فقط در یک کلاس رخ داده و ممکن است سیگنال واقعی باشد نه خطای ثبت؛ (ج) خطر اصلی، نشت بین Train و Test است، پس اثرش اندازه‌گیری شد نه حدس زده: فقط {twins} ردیف Test ({twins/n_te:.2%}) یک همزاد دقیق در Train دارند و Accuracy برای K=5 با و بدون آنها 0.9166 و 0.9167 است. اگر مدل نهایی در محیط واقعی ساخته شود، حذف تکراری‌ها پیش از تقسیم گزینهٔ محافظه‌کارانه‌تر است، ولی برای این تکلیف تغییر نتیجه‌ای ندارد.")
o3 = iqr.head(3)
B(f"**مقادیر پرت (IQR):** بیشترین سهم در {o3.index[0]} ({o3.Percent.iloc[0]}%)، {o3.index[1]} ({o3.Percent.iloc[1]}%) و {o3.index[2]} ({o3.Percent.iloc[2]}%). Area حدود {iqr.loc['Area','Percent']}% (عمدتاً دانه‌های درشت Bombay). این‌ها خطا نیستند و حذف نشدند، ولی KNN و K-Means به آن‌ها حساس‌اند.")

H("3. توزیع کلاس‌ها و عدم توازن")
P(f"کلاس‌ها متوازن نیستند: بزرگ‌ترین کلاس {big.title()} ({int(cd.Count.iloc[0]):,} نمونه، {cd.Percent.iloc[0]}%) و کوچک‌ترین {small.title()} ({int(cd.Count.iloc[-1]):,} نمونه، {cd.Percent.iloc[-1]}%) است؛ نسبت **{ratio}** به 1. مدلی که همه‌چیز را {big.title()} بگوید بدون یادگیری حدود {cd.Percent.iloc[0]:.0f}% Accuracy دارد. چون Accuracy کلاس‌های بزرگ را بیشتر وزن می‌دهد، خطای کامل روی یک کلاس کوچک در آن تقریباً دیده نمی‌شود؛ پس Recall و F1 هر کلاس و میانگین Macro هم گزارش شد.")

H("4. تحلیل اکتشافی و آماری")
P(f"میانگین Area کلاس‌ها از حدود {cm_means.Area.min()/1000:.0f} هزار (Dermason) تا {cm_means.Area.max()/1000:.0f} هزار (Bombay) تغییر می‌کند (نسبت {cm_means.Area.max()/cm_means.Area.min():.1f})، اما میانگین Roundness و Compactness فقط {cm_means.Roundness.max()/cm_means.Roundness.min():.2f} و {cm_means.Compactness.max()/cm_means.Compactness.min():.2f} برابر تفاوت دارد؛ یعنی اندازه Bombay را جدا می‌کند و بقیهٔ کلاس‌ها را بیشتر شکل از هم متمایز می‌کند. توزیع Area چوله به راست است (چولگی {df['Area'].skew():.2f}؛ میانگین {desc.loc['Mean','Area']:,.0f} بیشتر از میانه {desc.loc['Median','Area']:,.0f}). Bombay در Boxplot کاملاً جداست و Dermason، Seker و Sira هم‌پوشانی دارند. جدول 1 آمار توصیفی هفت ویژگی اصلی است (P25/P50/P75 و میانگین هر کلاس در tables/ ذخیره شده‌اند).", keep=False)
rows = []
for f_ in desc.columns:
    d = desc[f_]; fm = (lambda v: f"{v:,.0f}") if d["Mean"] > 100 else (lambda v: f"{v:.3f}")
    rows.append([f_, fm(d["Mean"]), fm(d["Median"]), fm(d["Std"]), fm(d["Min"]), fm(d["Max"]), fm(d["P90"])])
table(rows, ["Feature", "Mean", "Median", "Std", "Min", "Max", "P90"], [3.6, 2.2, 2.2, 2.2, 2.2, 2.2, 2.2])
fig(["class_distribution.png", "area_boxplot_by_class.png"], [8.2, 8.2], "شکل 1 (راست): تعداد و درصد هر کلاس؛ شکل 2 (چپ): Boxplot ویژگی Area به تفکیک کلاس")

H("5. تقسیم داده، Scale و اعتبار روش (بدون نشت)")
P(f"تقسیم با train_test_split و test_size=0.2، random_state=42 و stratify=y انجام شد: Train = {n_tr:,} و Test = {n_te:,} نمونه؛ درصد هر کلاس در Train و Test با درصد کل حداکثر 0.05 واحد تفاوت دارد. StandardScaler فقط روی Train آموزش دید (fit) و روی Train و Test فقط transform شد؛ میانگین و انحراف معیار Train پس از Scale دقیقاً 0 و 1 است و برای Test نزدیک آنها (مثلاً میانگین Area در Test برابر {sc.loc['Area','Mean (scaled test)']:.3f}). در Cross-Validation هم Scaler درون Pipeline و در هر Fold دوباره آموزش می‌بیند. نوت‌بوک 12 آزمون صریح نشت (شامل بازتولید تقسیم، تفاوت میانگین Scaler با میانگین کل داده، و fit فقط روی Train برای KNN و K-Means) را اجرا می‌کند و همه قبول شدند.")
P(f"**چرا Scale؟** KNN و K-Means بر فاصلهٔ اقلیدسی تکیه دارند. در داده‌های خام، Area و ConvexArea به‌تنهایی حدود {sh_area+sh_conv:.0f}% ({sh_area:.0f}% + {sh_conv:.0f}%) مربع فاصلهٔ اقلیدسی بین دو دانه را می‌سازند و Roundness یا Solidity عملاً نادیده گرفته می‌شوند. اثر عملی: Test Accuracy در KNN (K=5) از {k['raw']._5:.3f} (بدون Scale) به {acc[5]:.3f} (با Scale) و ARI در K-Means از {km['ARI_test_unscaled']:.2f} به {km['ARI_test']:.2f} می‌رسد. Scale مشکل واحدها و دامنه‌ها را حل می‌کند، نه هم‌بستگی و افزونگی ویژگی‌ها را (بخش 10).")

H("6. روش و نتایج KNN")
P("KNeighborsClassifier با فاصلهٔ اقلیدسی پیش‌فرض و K = 3، 5 و 9 روی ویژگی‌های Scale‌شدهٔ Train آموزش دید و روی Test سنجیده شد. ردیف آخر جدول 2 فقط مرجع (بدون Scale) است.", keep=True)
rows = []
for r in knn.itertuples():
    rows.append(["KNN", r.K, r.Scale, f4(r._4), f4(r._5), f4(r._6), f4(r._7)])
table(rows, ["Model", "K", "Scaling", "Train Acc", "Test Acc", "Macro Recall", "Macro F1"], [2, 1.2, 3.2, 2.4, 2.4, 2.8, 2.4], hl_rows=(3,))
P(f"جدول 2: نتایج KNN. K=3، 5 و 9 روی Test تقریباً یکسان‌اند: اختلاف Accuracy حدود {spread:.4f} (حدود {round(spread*n_te)} نمونه از {n_te:,}) است، در حالی که خطای معیار Accuracy حدود {se:.4f} (بازهٔ 95% حدود ±{1.96*se:.3f}) است. پس از این تفاوت‌ها **نمی‌توان «بهترین K عمومی» نتیجه گرفت**. Train Accuracy با بزرگ‌تر شدن K کاهش می‌یابد (از {k[3]._4:.3f} به {k[9]._4:.3f}) چون مدل هموارتر می‌شود؛ این روی تعمیم‌پذیری Test اثری نشان نداد. برای انتخاب حرفه‌ای K، 5-Fold Cross-Validation فقط روی Train (با Scaler درون Fold) اجرا شد: K=1 به‌وضوح ضعیف‌تر است (Macro-F1 برابر {cv.loc[1,'CV macro-F1 (mean)']:.4f})، از K=5 تا 21 منحنی تقریباً تخت است (بیشینه در K=9 با {cv.loc[9,'CV macro-F1 (mean)']:.4f} و K=11 با {cv.loc[11,'CV macro-F1 (mean)']:.4f}) و اختلاف‌ها هم‌مرتبهٔ انحراف معیار Foldها ({cv['CV macro-F1 (std)'].loc[[5,9,11]].min():.3f} تا {cv['CV macro-F1 (std)'].loc[[5,9,11]].max():.3f}) است.", size=9.5)
rows = [[c.title(), f3(r.Precision), f3(r.Recall), f3(r["F1-score"]), int(r.Support)] for c, r in pc5.iterrows()]
low_idx = [n for n, c in enumerate(pc5.index) if c == low]
table(rows, ["Class (K=5, Test)", "Precision", "Recall", "F1", "Support"], [4, 2.4, 2.4, 2.4, 2.4], hl_rows=low_idx)

H("7. ماتریس اغتشاش هفت‌کلاسه")
fig(["confusion_matrix_knn_k5.png"], [9.6], "شکل 3: ماتریس اغتشاش KNN با K=5 روی Test (اعداد و درصد سطری؛ ماتریس‌های K=3 و K=9 در figures/ و tables/ ذخیره شده‌اند)")
hi = classes[int(np.argmax(diag))]
B(f"**بیشترین پیش‌بینی صحیح:** {hi.title()} با {int(diag.max())} نمونه؛ عمدتاً چون بزرگ‌ترین کلاس است. **کمترین Recall:** {low.title()} با {rec5[low]:.3f}، و بلافاصله {low2.title()} با {rec5[low2]:.3f} (در K=9 ترتیب این دو عوض می‌شود، پس در حد نوسان برابرند). Bombay با Recall و Precision برابر 1.0 کاملاً جداست.")
B(f"**بیشترین اشتباه:** {pair[0].title()} ↔ {pair[1].title()}: {ab} خطا از {pair[0].title()} به {pair[1].title()} و {ba} خطا از {pair[1].title()} به {pair[0].title()}، جمعاً **{ab+ba}** خطا. جفت بعدی Barbunya ↔ Cali با {c2} خطاست.")
s_ = sim
B(f"**دلیل بر اساس ویژگی‌ها:** شکل Dermason و Sira شبیه است: تفاوت میانگین Extent ({s_.loc['Extent'].iloc[2]:.2f}) و Solidity ({s_.loc['Solidity'].iloc[2]:.2f}) کمتر از 0.1 انحراف معیار (Cohen d) و برای Eccentricity، Compactness، Roundness، AspectRatio و ShapeFactor3 کمتر از 0.9 است. تفاوت اصلی اندازه است (Area حدود {s_.loc['Area'].iloc[0]/1000:.0f} هزار در برابر {s_.loc['Area'].iloc[1]/1000:.0f} هزار؛ d = {s_.loc['Area'].iloc[2]:.1f})، ولی دامنهٔ Sira وسیع است و در ناحیهٔ مرزی با Dermason هم‌پوشانی می‌کند. Barbunya و Cali هر دو دانه‌های درشت‌اند (Area حدود {cm_means.Area['BARBUNYA']/1000:.0f} و {cm_means.Area['CALI']/1000:.0f} هزار).")

H("8. روش و نتایج K-Means")
P("KMeans با n_clusters=7، random_state=42 و n_init=10 فقط روی 16 ویژگی عددی Scale‌شدهٔ Train آموزش دید؛ ستون Class هرگز استفاده نشد. سپس Cluster نمونه‌های Test با predict تعیین شد. شمارهٔ Cluster فقط یک شناسه است، پس **هیچ Accuracy مستقیمی بین شمارهٔ Cluster و Class محاسبه نشد**؛ به‌جای آن از Crosstab، ARI، NMI و Purity (نگاشت مستند اکثریت) استفاده شد. جدول 3 Crosstab روی Test است.", keep=True)
rows = [[c.title()] + [int(v) for v in ct.loc[c]] for c in ct.index]
mx = {(ct.index.get_loc(c), int(list(ct.columns).index(ct.loc[c].idxmax()))+1) for c in ct.index}
cols = list(ct.columns)
table(rows, ["Class \\ Cluster"] + [str(c) for c in cols], [3.4] + [1.7]*len(cols))
rows = []
for c_, r in pur.iterrows():
    other = ktab.loc[c_, "Other important classes (>=5%)"]
    rows.append([c_, r["Dominant class (train)"], f"{r['Train size']:,}", f"{r['Train purity %']}%", int(r["Test size"]), f"{r['Test purity %']}%", other])
table(rows, ["Cluster", "Dominant class", "Train", "Train purity", "Test", "Test purity", "Other classes (Train, ≥5%)"], [1.9, 2.4, 1.4, 2.0, 1.3, 2.0, 5.2], size=8)
P(f"جدول 4: خلوص هر Cluster. **Purity کلی:** {km['purity_train']:.1%} روی Train و {km['purity_test']:.1%} روی Test (نگاشت اکثریتِ گرفته‌شده از Train روی Test نیز {km['test_accuracy_of_train_majority_mapping']:.1%} می‌دهد). این عدد با ARI={km['ARI_test']:.3f}، NMI={km['NMI_test']:.3f} و Silhouette={km['silhouette_train']:.2f} (Train) خوانده شود. Purity با افزایش تعداد خوشه‌ها بالا می‌رود و نگاشت اکثریت چندبه‌یک است: **Barbunya در هیچ Cluster غالب نیست** و با این نگاشت هرگز پیش‌بینی نمی‌شود.", size=9)
fig(["class_cluster_comparison.png"], [15.8], "شکل 4: Eccentricity در برابر Area با محورهای یکسان؛ چپ: کلاس واقعی، راست: Cluster (Test). نمودار جفت Compactness/Roundness در figures/ است.")
fig(["class_cluster_crosstab_heatmap.png", "knn_cv_k_selection.png"], [8.4, 7.6], "شکل 5 (راست): هیت‌مپ Crosstab کلاس–Cluster روی Test (رنگ = درصد کلاس در Cluster، عدد = تعداد)؛ شکل 6 (چپ): انتخاب K با Cross-Validation روی Train")
B("**انطباق تقریبی** (درصدها از Train): Cluster 3 دقیقاً Bombay (100%)، Cluster 0 عمدتاً Seker (92%)، Cluster 2 عمدتاً Dermason (92%) و Cluster 5 عمدتاً Horoz (94%).")
B("**یک کلاس در چند Cluster:** حدود 13% Dermason در Train (15% در Test) در Cluster 4؛ حدود 15% Cali و 9 تا 10% Horoz در Cluster 6؛ Barbunya بین Clusterهای 1، 4 و 6 پخش است.")
B("**یک Cluster با چند کلاس:** Cluster 1 نیمی Cali (53%) و نیمی Barbunya (45%)؛ Cluster 6 ترکیب Cali (49%) و Horoz (36%)؛ Cluster 4 (Sira با 75%) حدود 15% Dermason دارد.")
B(f"**چرا 7 خوشه ≠ کشف دقیق 7 کلاس؟** n_clusters=7 فقط تعداد ناحیه‌ها را تعیین می‌کند؛ K-Means گروه‌های گرد و فشرده می‌سازد، در حالی که کلاس‌ها هم‌پوشان (Sira/Dermason، Barbunya/Cali)، کشیده (Horoz) یا چندزیرگروه‌اند. بیشترین Silhouette در k=3 ({kscan.loc[3,'Silhouette']:.2f}) است و k=7 مقدار {kscan.loc[7,'Silhouette']:.2f} دارد؛ خود داده هفت گروه کاملاً جدا پیشنهاد نمی‌کند. با random_state=7 گروه‌بندی تقریباً یکسان (ARI برابر 0.999) ولی شماره‌ها متفاوت است؛ پس شناسهٔ Cluster معنای کلاس ندارد.")

H("9. مقایسهٔ KNN و K-Means")
P(f"KNN مدل **نظارت‌شده** است، برچسب را می‌بیند و خروجی‌اش هم‌جنس برچسب است؛ پس Accuracy، Precision، Recall و F1 دارد. K-Means **بدون نظارت** است و فقط شناسهٔ گروه می‌دهد که معنی ذاتی ندارد. مقایسه با یک عدد Accuracy نادرست است: هر مدل با معیار مناسب خودش (KNN: Macro Recall/F1 و ماتریس اغتشاش؛ K-Means: Crosstab، ARI، NMI، Purity) سنجیده شد. جمع‌بندی: با برچسب، KNN حدود {acc[5]:.0%} نمونه‌های Test را درست تشخیص می‌دهد (Macro F1 حدود {k[5]._7:.2f})؛ بدون برچسب، K-Means تا حدی (ARI برابر {km['ARI_test']:.2f}) ساختار کلاس‌ها را بازیابی می‌کند.")

H("10. چالش ابعاد، فاصله و افزونگی ویژگی‌ها")
B(f"**ضعف مفهوم فاصله:** با افزایش بُعد، فاصلهٔ نزدیک‌ترین و دورترین نقطه به هم نزدیک می‌شود. در شبیه‌سازی ما تفاوت نسبی فاصلهٔ دورترین و نزدیک‌ترین نقطه (نسبت به فاصلهٔ نزدیک‌ترین) از {dist.iloc[0,0]:.1f} در 2 بُعد به {dist.loc[16].iloc[0]:.1f} در 16 بُعد و {dist.loc[1024].iloc[0]:.2f} در 1024 بُعد می‌رسد. 16 بُعد هنوز قابل قبول است، ولی اثر شروع شده و «نزدیک‌ترین همسایه» اطلاعات کمتری دارد.")
B(f"**استقلال ویژگی‌ها:** 16 ویژگی مستقل نیستند؛ فقط {n90} جهت از 16 جهت 90% واریانس استانداردشده و {n99} جهت 99% آن را می‌سازد و {corr_pairs} جفت ویژگی |r| ≥ 0.95 دارند.")
B(f"**Area، Perimeter و EquivDiameter:** Area–Perimeter r={r_ap:.3f}، Area–EquivDiameter r={r_ae:.3f}، Perimeter–EquivDiameter r={r_pe:.3f} (ConvexArea با Area نیز r=0.9999). یک اطلاعات (اندازه) چند بار در فاصله شمرده می‌شود، پس وزن ضمنی اندازه از شکل بیشتر می‌شود. Compactness و Eccentricity هم r={r_ce:.2f} دارند.")
fig(["correlation_matrix.png"], [9.2], "شکل 7: ماتریس هم‌بستگی 16 ویژگی؛ بلوک بزرگ (Area، Perimeter، ConvexArea، EquivDiameter، MajorAxisLength، MinorAxisLength) و بلوک شکل (Compactness، Eccentricity، AspectRatio، ShapeFactor3) پرهم‌بستگی‌اند")
B("**Scale و ویژگی‌های تکراری:** Scale فقط دامنه‌ها را یکسان می‌کند و هم‌بستگی را تغییر نمی‌دهد؛ پس مشکل واحدها حل می‌شود ولی تکراری‌بودن اطلاعات می‌ماند. (طبق دستور پروژه PCA یا انتخاب ویژگی به کار نرفت و فقط تحلیل شد.)")

H("11. پاسخ پرسش‌های نهایی")
B(f"**1. متوازن؟** خیر؛ نسبت {ratio} به 1 ({big.title()} در برابر {small.title()}).")
B(f"**2. کمترین Recall؟** {low.title()} ({rec5[low]:.3f}) و پس از آن {low2.title()} ({rec5[low2]:.3f}).")
B(f"**3. بیشترین اشتباه؟** {pair[0].title()} و {pair[1].title()} ({ab+ba} خطا).")
B(f"**4. اثر K؟** برای K=3/5/9 بسیار کم (چند نمونه، کمتر از خطای معیار)؛ Train Accuracy با K کاهش می‌یابد؛ CV فقط نشان می‌دهد K=1 بدتر است و بازهٔ میانی هم‌ارز.")
B(f"**5. اهمیت Scale؟** بدون Scale، Area و ConvexArea فاصله را تسخیر می‌کنند: KNN از {k['raw']._5:.3f} به {acc[5]:.3f} و K-Means ARI از {km['ARI_test_unscaled']:.2f} به {km['ARI_test']:.2f} می‌رسد.")
B(f"**6. انطباق Clusterها با کلاس‌ها؟** جزئی (ARI {km['ARI_test']:.2f}، NMI {km['NMI_test']:.2f}، Purity {km['purity_test']:.0%}): Bombay کامل؛ Seker، Dermason و Horoz حدود 92 تا 94%؛ Barbunya و Cali در هم آمیخته‌اند.")
B("**7. شمارهٔ Cluster؟** شناسهٔ دلخواه است و با random_state یا اجرای دیگر عوض می‌شود؛ بدون نگاشت مستند معنای کلاس ندارد.")
B("**8. تفاوت طبقه‌بندی و خوشه‌بندی؟** در طبقه‌بندی نظارت‌شده مدل از جفت (ویژگی، برچسب) یاد می‌گیرد و برچسب پیش‌بینی می‌کند؛ در خوشه‌بندی بدون نظارت فقط ویژگی‌ها هست و الگوریتم گروه‌های مشابه می‌سازد.")
B("**9. محدودیت در خط بسته‌بندی؟** تغییر نور/دوربین/سرعت نوار (Data Drift)؛ دانه‌های چسبیده، شکسته یا ناخالصی و نوع ناشناخته (KNN همیشه یکی از 7 کلاس را می‌دهد)؛ حدود 8% خطا و Precision حدود 0.84 برای Sira؛ وابستگی Precision به ترکیب کلاس‌ها در خط؛ هزینهٔ محاسبهٔ فاصله با کل داده؛ نیاز به آزمون روی داده‌های واقعی در زمان‌های مختلف.")

H("12. نتیجه‌گیری و محدودیت‌ها")
P(f"**نتیجه:** ویژگی‌های هندسی به‌تنهایی برای تشخیص نوع لوبیا کافی‌اند: KNN با Scale روی Test حدود {acc[5]:.1%} Accuracy و Macro F1 حدود {k[5]._7:.3f} می‌دهد و Bombay را بی‌خطا جدا می‌کند؛ ضعیف‌ترین بخش، جفت Dermason–Sira (و سپس Barbunya–Cali) است. K-Means بدون برچسب فقط ساختار جزئی را کشف می‌کند (ARI {km['ARI_test']:.2f}، Purity {km['purity_test']:.0%}) و بدون Scale عملاً بی‌معناست.")
P("**محدودیت‌ها:** (1) همهٔ نتایج روی یک تقسیم Train/Test هستند و فاصلهٔ اطمینان حدود ±1% دارند؛ تفاوت‌های کوچک بین K یا بین مدل‌ها معنادار نیست. (2) 68 ردیف تکراری و مقادیر پرت عمداً حفظ شدند؛ اثر نشت ناشی از تکراری‌ها اندازه‌گیری و ناچیز بود، ولی حذف آن‌ها گزینهٔ دیگری است. (3) تحلیل اکتشافی (هم‌بستگی، میانگین کلاس‌ها) روی کل داده انجام شد و فقط توصیفی است، در آموزش هیچ مدلی نقش نداشت. (4) ویژگی‌ها از تصاویر یک سامانهٔ تصویربرداری استخراج شده‌اند؛ تعمیم به نور، دوربین یا نمونهٔ دیگر آزمون نشده است. (5) K-Means خوشه‌های گرد فرض می‌کند و Purity به تعداد خوشه‌ها و نگاشت اکثریت وابسته است. (6) PCA، انتخاب ویژگی و تنظیم فراپارامتر (وزن فاصله، معیار فاصله) طبق دستور پروژه انجام نشد و می‌تواند در کارهای بعدی بررسی شود.", size=9)

H("13. فایل‌های تحویلی و تکرارپذیری")
P("notebooks/Dry_Bean_KNN_KMeans.ipynb (اجرای کامل بدون خطا)، README.md، figures/ (13 نمودار از جمله ماتریس اغتشاش هفت‌کلاسه و دو نمودار مقایسهٔ Class/Cluster)، tables/ (جدول‌ها از جمله Crosstab، Purity و نتایج نهایی KNN و K-Means به‌همراه results_summary.json) و این گزارش. اجرای مجدد: Restart & Run All بدون تغییر در تنظیمات.", size=9)

doc.core_properties.title = "Dry Bean Final Report"; doc.core_properties.author = "Dry Bean Project"
out = ROOT/"reports"/"Dry_Bean_Final_Report.docx"; doc.save(out); print("saved", out)
